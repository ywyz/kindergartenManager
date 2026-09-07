# ruff: noqa: BLE001 - the released boundary sanitizes template/parser failures
"""Fixed-template local Word port for the released weekly/monthly pair.

The caller cannot choose a path, template version, or fallback.  Construction
loads and verifies the two release assets once; rendering then works only from
those verified bytes and returns an in-memory DOCX.
"""

from __future__ import annotations

import json
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from docx import Document

from app.service.template_center.contracts import (
    DocumentType,
    ExportParseReport,
    RenderedTemplate,
    TemplateExportBinding,
)
from app.service.template_center.registry import CANDIDATE_QUALIFICATION_PROFILES
from app.service.weekly_monthly_plans.contracts import (
    MonthlyThemeActivityPlan,
    WeeklyActivityPlan,
)
from app.service.weekly_monthly_plans.export_contracts import PlanDocumentType
from app.service.weekly_monthly_plans.formal_exporter import _payload_sha256

_ROOT = Path(__file__).resolve().parents[3]
_RELEASED = {
    DocumentType.WEEKLY_ACTIVITY_PLAN: (
        _ROOT / "templates" / "weekplan.docx",
        UUID("00000000-0000-0000-0000-000000000801"),
    ),
    DocumentType.MONTHLY_THEME_ACTIVITY_PLAN: (
        _ROOT / "templates" / "monthplan.docx",
        UUID("00000000-0000-0000-0000-000000000802"),
    ),
}
_ACTIVE_VERSION = 8


class ReleasedWordPortError(RuntimeError):
    """Stable, content-free local template failure."""

    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


def _kind(value: object) -> DocumentType:
    try:
        return value if type(value) is DocumentType else DocumentType(value)
    except (TypeError, ValueError):
        raise ReleasedWordPortError("document_type_invalid") from None


def _profile(document_type: DocumentType):
    matches = tuple(
        profile
        for profile in CANDIDATE_QUALIFICATION_PROFILES
        if profile.document_type is document_type and profile.profile_version == 2
    )
    if len(matches) != 1:
        raise ReleasedWordPortError("profile_invalid")
    return matches[0]


def _joined(values: tuple[str, ...]) -> str:
    return "\n".join(values)


def _fill_weekly(document: Document, plan: WeeklyActivityPlan) -> None:
    header = (
        f"主题名称：{plan.theme_name}    班级：{plan.scope.class_name}    "
        f"第（{plan.period.week_number}）周（{plan.period.week_start:%Y年%m月%d日}"
        f"——{plan.period.week_end:%Y年%m月%d日}）"
    )
    people = (
        f"教师：{'、'.join(plan.scope.teacher_names)}    "
        f"保育员：{plan.scope.caregiver_name or ''}"
    )
    for index in (0, 4):
        document.paragraphs[index].text = "幼儿园每周工作计划表"
        document.paragraphs[index + 1].text = header
        document.paragraphs[index + 2].text = people
    for table in document.tables:
        table.cell(0, 0).text = f"第{plan.period.week_number}周"
        for offset, day in enumerate(plan.days, start=2):
            table.cell(0, offset).text = f"{day.weekday_cn}\n{day.day_date:%m月%d日}"
            table.cell(1, offset).text = day.morning_talk
            table.cell(2, offset).text = day.collective_activity
        # The released template merges each game row across all five weekdays.
        # Write it once, retaining the day association and original body order.
        table.cell(3, 2).text = "\n".join(
            f"{day.weekday_cn}：{day.outdoor_game}" for day in plan.days
        )
        table.cell(4, 2).text = "\n".join(
            f"{day.weekday_cn}：{day.area_game}" for day in plan.days
        )
        for row, value in (
            (5, plan.weekly_focus),
            (6, plan.environment_creation),
            (7, plan.life_habits),
            (8, plan.home_school_cooperation),
        ):
            table.cell(row, 2).text = value


def _fill_monthly(document: Document, plan: MonthlyThemeActivityPlan) -> None:
    document.paragraphs[0].text = "幼儿园主题教育活动计划"
    document.paragraphs[1].text = (
        f"班级：{plan.scope.class_name}    执行年月："
        f"{plan.period.year:04d}.{plan.period.month:02d}    "
        f"带班老师：{'、'.join(plan.scope.teacher_names)}    "
        f"保育老师：{plan.scope.caregiver_name or ''}"
    )
    table = document.tables[0]
    table.cell(1, 0).text = (
        f"本月主题：{plan.theme_name}\n"
        f"上月分析：{plan.previous_month_analysis}\n"
        f"本月重点：{plan.monthly_focus}"
    )
    table.cell(3, 1).text = _joined(plan.theme_goals)
    table.cell(3, 3).text = _joined(plan.life_habits)
    table.cell(4, 1).text = _joined(plan.play_activities)
    table.cell(4, 3).text = _joined(plan.environment_creation)
    table.cell(5, 1).text = _joined(plan.home_school_cooperation)
    table.cell(5, 3).text = _joined(plan.other)
    table.cell(7, 0).text = _joined(plan.activity_contents)


class ReleasedWeeklyMonthlyWordPort:
    """Render and parse only the immutable release-qualified template bytes."""

    def __init__(self, templates: dict[DocumentType, bytes]) -> None:
        self._templates = dict(templates)
        self._payload_by_rendered_sha: dict[str, str] = {}

    async def resolve_active(
        self, tenant_id: int, document_type: object
    ) -> TemplateExportBinding:
        if type(tenant_id) is not int or tenant_id <= 0:
            raise ReleasedWordPortError("tenant_invalid")
        key = _kind(document_type)
        profile = _profile(key)
        _, template_version_id = _RELEASED[key]
        return TemplateExportBinding(
            document_type=key,
            content_sha256=profile.seed_handle.expected_sha256,
            contract_id=profile.contract.contract_id,
            contract_version=profile.contract.contract_version,
            tenant_id=tenant_id,
            template_version_id=template_version_id,
            version=_ACTIVE_VERSION,
        )

    async def render(
        self, binding: TemplateExportBinding, payload: object
    ) -> RenderedTemplate:
        if type(binding) is not TemplateExportBinding:
            raise ReleasedWordPortError("binding_invalid")
        key = binding.document_type
        profile = _profile(key)
        if (
            binding.content_sha256 != profile.seed_handle.expected_sha256
            or binding.contract_id != profile.contract.contract_id
            or binding.contract_version != profile.contract.contract_version
            or binding.version != _ACTIVE_VERSION
            or binding.template_version_id != _RELEASED[key][1]
        ):
            raise ReleasedWordPortError("binding_invalid")
        try:
            document = Document(BytesIO(self._templates[key]))
            public_type = PlanDocumentType(key.value)
            if key is DocumentType.WEEKLY_ACTIVITY_PLAN:
                if type(payload) is not WeeklyActivityPlan:
                    raise ReleasedWordPortError("payload_invalid")
                _fill_weekly(document, payload)
            else:
                if type(payload) is not MonthlyThemeActivityPlan:
                    raise ReleasedWordPortError("payload_invalid")
                _fill_monthly(document, payload)
            buffer = BytesIO()
            document.save(buffer)
            rendered_bytes = buffer.getvalue()
            rendered_sha = sha256(rendered_bytes).hexdigest()
            payload_sha = _payload_sha256(public_type, payload)
            self._payload_by_rendered_sha[rendered_sha] = payload_sha
            return RenderedTemplate(
                binding=binding,
                rendered_bytes=rendered_bytes,
                rendered_sha256=rendered_sha,
                payload_sha256=payload_sha,
            )
        except ReleasedWordPortError:
            raise
        except BaseException:
            raise ReleasedWordPortError("render_failed") from None

    async def parse(
        self, binding: TemplateExportBinding, rendered_bytes: bytes
    ) -> ExportParseReport:
        if (
            type(binding) is not TemplateExportBinding
            or type(rendered_bytes) is not bytes
        ):
            raise ReleasedWordPortError("parse_input_invalid")
        rendered_sha = sha256(rendered_bytes).hexdigest()
        payload_sha = self._payload_by_rendered_sha.pop(rendered_sha, None)
        if payload_sha is None:
            raise ReleasedWordPortError("render_identity_unknown")
        try:
            with ZipFile(BytesIO(rendered_bytes)) as package:
                names = tuple(sorted(package.namelist()))
                if "word/document.xml" not in names:
                    raise ReleasedWordPortError("package_invalid")
                forbidden = tuple(
                    name
                    for name in names
                    if name.casefold().endswith(".bin")
                    or "/activex/" in name.casefold()
                    or "/embeddings/" in name.casefold()
                    or "/oleobject" in name.casefold()
                )
                external = tuple(
                    name
                    for name in names
                    if name.endswith(".rels")
                    and b'TargetMode="External"' in package.read(name)
                )
                structure_sha = sha256(
                    json.dumps(names, separators=(",", ":")).encode("utf-8")
                ).hexdigest()
        except ReleasedWordPortError:
            raise
        except (BadZipFile, KeyError, OSError):
            raise ReleasedWordPortError("package_invalid") from None
        return ExportParseReport(
            binding=binding,
            valid=not forbidden and not external,
            structure_summary_sha256=structure_sha,
            unresolved_token_ids=(),
            has_macros=bool(forbidden),
            has_external_relationships=bool(external),
            rendered_sha256=rendered_sha,
            payload_sha256=payload_sha,
        )


def build_released_weekly_monthly_word_port() -> ReleasedWeeklyMonthlyWordPort:
    templates: dict[DocumentType, bytes] = {}
    try:
        for document_type, (path, _) in _RELEASED.items():
            content = path.read_bytes()
            expected = _profile(document_type).seed_handle.expected_sha256
            if sha256(content).hexdigest() != expected:
                raise ReleasedWordPortError("template_identity_mismatch")
            templates[document_type] = content
    except ReleasedWordPortError:
        raise
    except BaseException:
        raise ReleasedWordPortError("template_unavailable") from None
    return ReleasedWeeklyMonthlyWordPort(templates)


__all__ = (
    "ReleasedWeeklyMonthlyWordPort",
    "ReleasedWordPortError",
    "build_released_weekly_monthly_word_port",
)
