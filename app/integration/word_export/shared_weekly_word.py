"""Independent v3 layout qualification; no legacy qualification inheritance.

Local LibreOffice inspection never activates a template or certifies Word.
Until a independently qualified authority is installed, formal resolution fails.
"""

from __future__ import annotations

import asyncio
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from collections import Counter
from contextlib import asynccontextmanager
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from app.service.academic_identity.contracts import IdentityRejected
from app.service.shared_weekly.authoring_contracts import (
    AREA_TITLE,
    OUTDOOR_TITLE,
    WeeklyAuthoringDraft,
)
from app.service.shared_weekly.layout_authority import LayoutAuthority
from app.service.shared_weekly.layout_contracts import (
    LayoutBinding,
    RenderedWeek,
    WeekDisplay,
)

SEED_SHA256 = "f6c17c137f04e29a68524ed400eb395984e93a16c234a065b5794d9f49a9347b"
SEED_PATH = Path(__file__).resolve().parents[3] / "templates/weekplan.docx"


class LayoutRejected(ValueError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


def _complete(body):
    if type(body) is not WeeklyAuthoringDraft or body.calendar is None:
        raise LayoutRejected("schema_invalid")
    try:
        body.validate_complete()
    except IdentityRejected:
        raise LayoutRejected("required_fields_missing") from None
    if (
        not body.theme.strip()
        or not body.people.teachers
        or any(
            not value.strip()
            for value in (*body.people.teachers, body.people.caregiver)
        )
    ):
        raise LayoutRejected("required_fields_missing")
    for day, teaching, label in body.calendar.columns:
        item = next(d for d in body.days if d.day == day)
        if teaching and any(
            not getattr(item, field).strip()
            for field in (
                "morning_talk_topic",
                "morning_talk_questions",
                "activity_name",
            )
        ):
            raise LayoutRejected("required_fields_missing")


def _numbered(body, prefix):
    return " ".join(f"{i + 1}.{body.value_at(f'{prefix}.{i}')}" for i in range(3))


def fill_document(
    seed: bytes, body: WeeklyAuthoringDraft, display: WeekDisplay
) -> bytes:
    """Derive one v3 table from the exact controlled seed, preserving its style.

    The duplicate sample sheet is removed; this is a new profile requiring its
    own qualification, never a claim that the old seed already qualifies v3.
    """
    if sha256(seed).hexdigest() != SEED_SHA256:
        raise LayoutRejected("template_changed")
    if type(display) is not WeekDisplay:
        raise LayoutRejected("display_invalid")
    _complete(body)
    document = Document(BytesIO(seed))
    table_style = document.tables[0].style
    cell_borders = deepcopy(
        document.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:tcBorders"))
    )
    for child in list(document.element.body):
        if child.tag != qn("w:sectPr"):
            document.element.body.remove(child)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    # Retain controlled template margins, not historical compact trial margins.
    document.add_paragraph("幼儿园每周工作计划表")
    theme = body.theme.strip().strip("《》")
    document.add_paragraph(
        f"主题名称：《{theme}》    班级：{display.class_name}    {display.term_name} 第{display.week_number}周"
    )
    document.add_paragraph(
        f"教师：{'、'.join(body.people.teachers)}    保育员：{body.people.caregiver}"
    )
    count = len(body.days)
    table = document.add_table(rows=9, cols=count + 2)
    table.style = table_style
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell._tc.get_or_add_tcPr().append(deepcopy(cell_borders))
    width = section.page_width - section.left_margin - section.right_margin
    for col in table.columns:
        col.width = int(width / (count + 2))
    table.cell(0, 0).merge(table.cell(0, 1)).text = f"第{display.week_number}周"
    for offset, (day, teaching, label) in enumerate(body.calendar.columns, 2):
        table.cell(
            0, offset
        ).text = f"周{'一二三四五六日'[day.weekday()]}\n{day:%m月%d日}"
        item = body.days[offset - 2]
        table.cell(1, offset).text = (
            f"{item.morning_talk_topic}\n{item.morning_talk_questions}"
            if teaching
            else label
        )
        table.cell(2, offset).text = item.activity_name if teaching else label
    table.cell(1, 0).merge(table.cell(2, 0)).text = "学习活动"
    table.cell(1, 1).text = "晨间谈话"
    table.cell(2, 1).text = "集体活动"
    table.cell(3, 0).merge(table.cell(4, 0)).text = "游戏活动"
    table.cell(3, 1).text = "户外游戏"
    table.cell(4, 1).text = "区域游戏"
    games = [OUTDOOR_TITLE]
    for prefix in ("games.collective.0", "games.collective.1", "games.autonomous"):
        games.append(
            body.value_at(prefix + ".name") + "：" + _numbered(body, prefix + ".goals")
        )
    values = [
        "\n".join(games),
        f"{AREA_TITLE}\n重点区域：{body.value_at('area.name')}\n目标：{_numbered(body, 'area.goals')}\n材料：{body.value_at('area.materials')}\n指导：{_numbered(body, 'area.guidance')}",
        _numbered(body, "focus"),
        _numbered(body, "environment"),
        " ".join(
            f"{i + 1}.{body.value_at(f'habits.{i}.name')}：{body.value_at(f'habits.{i}.content')}"
            for i in range(3)
        ),
        body.value_at("home"),
    ]
    for row, value in enumerate(values, 3):
        table.cell(row, 2).merge(table.cell(row, count + 1)).text = value
    for row, label in enumerate(
        ("本周重点", "环境创设", "生活习惯培养", "家园共育"), 5
    ):
        table.cell(row, 0).merge(table.cell(row, 1)).text = label
    paragraphs = list(document.paragraphs)
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc not in seen:
                seen.add(cell._tc)
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(20)
        fmt.keep_with_next = False
        for run in paragraph.runs:
            run.font.name = "SimSun"
            run.font.size = Pt(12)
            fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                run._element.get_or_add_rPr().append(fonts)
            fonts.set(qn("w:eastAsia"), "宋体")
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    # Round-trip every visible value, including newline boundaries, in order.
    parsed = Document(BytesIO(data))
    expected = [p.text for p in document.paragraphs] + [
        c.text for r in table.rows for c in r.cells
    ]
    actual = [p.text for p in parsed.paragraphs] + [
        c.text for r in parsed.tables[0].rows for c in r.cells
    ]
    if actual != expected or len(parsed.tables) != 1:
        raise LayoutRejected("roundtrip_failed")
    return data


def _text_complete(document, xml, grid):
    """Check rendered glyph counts and cell-local reading order.

    PDF layout extraction interleaves columns at every physical line. Restrict
    the actual word boxes to each cell's horizontal span before joining wrapped
    lines, so a legal wrapped date/holiday is not mistaken for missing text.
    """
    words = list(xml.iter("{http://www.w3.org/1999/xhtml}word"))
    if not words:
        return False
    lines = []
    for word in sorted(words, key=lambda word: float(word.attrib["yMin"])):
        top = float(word.attrib["yMin"])
        if not lines or top - lines[-1][0] > 3:
            lines.append((top, [word]))
        else:
            lines[-1][1].append(word)
    words = [
        word
        for _, line in lines
        for word in sorted(line, key=lambda word: float(word.attrib["xMin"]))
    ]

    clean = lambda text: "".join(text.split())
    full = clean("".join(word.text or "" for word in words))
    expected = [p.text for p in document.paragraphs]
    if any(clean(p.text) not in full for p in document.paragraphs if p.text):
        return False
    xs, ys = grid
    for table in document.tables:
        if len(xs) != len(table.columns) + 1 or len(ys) != len(table.rows) + 1:
            return False
        seen = set()
        for row_index, row in enumerate(table.rows):
            for index, cell in enumerate(row.cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                expected.append(cell.text)
                end_row = row_index + 1
                while (
                    end_row < len(table.rows)
                    and table.rows[end_row].cells[index]._tc is cell._tc
                ):
                    end_row += 1
                left, right = xs[index], xs[index + cell._tc.grid_span]
                top, bottom = ys[row_index], ys[end_row]
                selected = []
                for word in words:
                    a = word.attrib
                    x0, x1, y0, y1 = (
                        float(a[k]) for k in ("xMin", "xMax", "yMin", "yMax")
                    )
                    if (
                        left <= (x0 + x1) / 2 <= right
                        and top <= (y0 + y1) / 2 <= bottom
                    ):
                        if (
                            x0 < left - 0.75
                            or x1 > right + 0.75
                            or y0 < top - 0.75
                            or y1 > bottom + 0.75
                        ):
                            return False
                        selected.append(word)
                for position, word in enumerate(selected):
                    a = word.attrib
                    for other in selected[position + 1 :]:
                        b = other.attrib
                        overlap_x = min(float(a["xMax"]), float(b["xMax"])) - max(
                            float(a["xMin"]), float(b["xMin"])
                        )
                        overlap_y = min(float(a["yMax"]), float(b["yMax"])) - max(
                            float(a["yMin"]), float(b["yMin"])
                        )
                        if overlap_x > 0.5 and overlap_y > 1:
                            return False
                local = clean("".join(word.text or "" for word in selected))
                if clean(cell.text) != local:
                    return False
    # Ensure duplicate values are preserved too; finding one repeated label
    # elsewhere is not enough to prove every rendered occurrence survived.
    return Counter(clean("".join(expected))) == Counter(full)


def _rendered_grid(svg):
    """Read actual stroked borders from Poppler's renderer output, not OOXML."""
    horizontal, vertical = [], []
    number = r"(-?[0-9]+(?:\.[0-9]+)?)"
    for path in svg.iter("{http://www.w3.org/2000/svg}path"):
        if path.get("stroke") is None or path.get("fill") != "none":
            continue
        match = re.fullmatch(
            r"M\s+"
            + number
            + r"\s+"
            + number
            + r"\s+L\s+"
            + number
            + r"\s+"
            + number
            + r"\s*",
            path.get("d", ""),
        )
        transform = re.fullmatch(r"matrix\(([^)]+)\)", path.get("transform", ""))
        if match is None or transform is None:
            raise LayoutRejected("layout_geometry_invalid")
        a, b, c, d, e, f = (float(v.strip()) for v in transform[1].split(","))
        x0, y0, x1, y1 = map(float, match.groups())
        x0, y0, x1, y1 = (
            a * x0 + c * y0 + e,
            b * x0 + d * y0 + f,
            a * x1 + c * y1 + e,
            b * x1 + d * y1 + f,
        )
        if abs(y0 - y1) < 0.1 and abs(x1 - x0) > 5:
            horizontal.append(y0)
        elif abs(x0 - x1) < 0.1 and abs(y1 - y0) > 5:
            vertical.append(x0)
        else:
            raise LayoutRejected("layout_geometry_invalid")

    def clustered(values):
        groups = []
        for value in sorted(values):
            if not groups or value - groups[-1][-1] > 1:
                groups.append([value])
            else:
                groups[-1].append(value)
        return tuple(sum(group) / len(group) for group in groups)

    return clustered(vertical), clustered(horizontal)


async def _process(*args):
    process = await asyncio.create_subprocess_exec(
        *args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE
    )
    try:
        stdout, _ = await asyncio.wait_for(process.communicate(), 45)
    except BaseException:
        if process.returncode is None:
            process.kill()
        await process.wait()
        raise
    if process.returncode:
        raise LayoutRejected("renderer_failed")
    return stdout


class SharedWeeklyWordPort:
    """Fail-closed production boundary plus a separate local qualification probe."""

    def __init__(self, authority: LayoutAuthority | None = None):
        self.authority = authority or LayoutAuthority()

    async def resolve_binding(self, tenant_id: int) -> LayoutBinding:
        return await self.authority.resolve_binding(tenant_id)

    @asynccontextmanager
    async def binding_guard(self, binding: LayoutBinding):
        async with self.authority.binding_guard(binding):
            if sha256(SEED_PATH.read_bytes()).hexdigest() != binding.template_sha256:
                raise LayoutRejected("template_changed")
            yield

    async def render_check(self, binding, body, display):
        renderer = await self.authority.renderer(binding)
        if renderer["product"] != "LibreOffice":
            raise LayoutRejected("renderer_unavailable")
        if shutil.which("libreoffice") is None:
            raise LayoutRejected("renderer_missing")
        current = (await _process("libreoffice", "--version")).decode().strip()
        if current != renderer["version"]:
            raise LayoutRejected("renderer_changed")
        if binding.template_sha256 != SEED_SHA256:
            raise LayoutRejected("template_changed")
        result = await self._render(binding, body, display)
        async with self.binding_guard(binding):
            return result

    async def inspect_local(
        self, body: WeeklyAuthoringDraft, display: WeekDisplay
    ) -> RenderedWeek:
        """Actual LO report; local-unqualified is never a formal binding."""
        return await self._render(
            LayoutBinding(1, SEED_SHA256, "local-unqualified"), body, display
        )

    async def _render(self, binding, body, display):
        payload_hash = sha256(body.serialize().encode()).hexdigest()

        def result(fits=False, pages=0, reason="renderer_failed", data=None):
            return RenderedWeek(binding, payload_hash, fits, pages, reason, data)

        try:
            data = fill_document(SEED_PATH.read_bytes(), body, display)
            if any(
                shutil.which(name) is None
                for name in (
                    "libreoffice",
                    "pdftotext",
                    "pdfinfo",
                    "fc-match",
                    "pdftocairo",
                )
            ):
                return result(reason="renderer_missing")
            font = (await _process("fc-match", "-f", "%{family}", "SimSun")).decode()
            if "SimSun" not in font.split(","):
                return result(reason="font_missing")
            with tempfile.TemporaryDirectory(prefix="shared-weekly-layout-") as temp:
                root = Path(temp)
                source = root / "week.docx"
                source.write_bytes(data)
                await _process(
                    "libreoffice",
                    "-env:UserInstallation=" + (root / "profile").as_uri(),
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    temp,
                    str(source),
                )
                pdf = root / "week.pdf"
                info = (await _process("pdfinfo", str(pdf))).decode()
                pages = int(
                    next(
                        line.split(":")[1]
                        for line in info.splitlines()
                        if line.startswith("Pages:")
                    )
                )
                if pages != 1:
                    return result(pages=pages, reason="layout_overflow")
                await _process("pdftotext", "-bbox", str(pdf), str(root / "bbox.html"))
                xml = ET.parse(root / "bbox.html")
                for page in xml.iter("{http://www.w3.org/1999/xhtml}page"):
                    width, height = (
                        float(page.attrib["width"]),
                        float(page.attrib["height"]),
                    )
                    for word in page.iter("{http://www.w3.org/1999/xhtml}word"):
                        a = word.attrib
                        if (
                            float(a["xMax"]) <= float(a["xMin"])
                            or float(a["yMax"]) <= float(a["yMin"])
                            or float(a["xMin"]) < 0
                            or float(a["yMin"]) < 0
                            or float(a["xMax"]) > width
                            or float(a["yMax"]) > height
                        ):
                            return result(pages=pages, reason="layout_overflow")
                await _process("pdftotext", "-layout", str(pdf), str(root / "text.txt"))
                document = Document(BytesIO(data))
                await _process("pdftocairo", "-svg", str(pdf), str(root / "grid.svg"))
                grid = _rendered_grid(ET.parse(root / "grid.svg"))
                if not _text_complete(document, xml, grid):
                    return result(pages=pages, reason="layout_overflow")
                return result(True, pages, "fits", data)

        except (ValueError, OSError, StopIteration, ET.ParseError) as exc:
            return result(
                reason=exc.code if isinstance(exc, LayoutRejected) else "layout_invalid"
            )
