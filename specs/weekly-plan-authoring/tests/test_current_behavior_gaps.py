"""WP-A current-behavior probes, NOT a request to change legacy semantics.

These exercise existing production code with synthetic SQLite data. Failing
expectations express #77 requirements that cannot be obtained by reusing the
old contracts unchanged. GREEN must eventually target the new shared contract;
do not weaken legacy authorization or alter monthly behavior to pass probes.
No missing imports, xfail, fixed failures, real AI, or business databases.
"""

from dataclasses import replace
from datetime import UTC, date, datetime
from io import BytesIO

import pytest
import pytest_asyncio
from docx import Document
from sqlalchemy import event, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app.core.models  # noqa: F401
from app.core.database import Base
from app.core.models.user import User, UserRole
from app.core.models.weekly_monthly_plan import (
    WeeklyMonthlyPlan,
    WeeklyMonthlyPlanVersion,
)
from app.integration.ai_client import lesson_plan_client
from app.integration.word_export.released_weekly_monthly_word_port import (
    build_released_weekly_monthly_word_port,
)
from app.repository.daily_plan_repository import (
    get_daily_plan_by_id_for_user,
    list_daily_plans_for_user,
    save_daily_plan,
)
from app.repository.weekly_monthly_plan_repository import (
    SqlAlchemyPlanAggregateReadRepository,
)
from app.service.date_service import get_week_number
from app.service.weekly_monthly_plans.authorization import (
    DatabasePlanAuthorizationAdapter,
)
from app.service.weekly_monthly_plans.contracts import (
    PlanAction,
    PlanAuthorizationRequest,
    PlanKind,
    ReviewStatus,
    WeeklyDay,
)
from app.service.weekly_monthly_plans.qualification_orchestration import (
    _canonical_weekly_plan,
)


@pytest_asyncio.fixture
async def db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")

    @event.listens_for(engine.sync_engine, "connect")
    def enable_foreign_keys(connection, _record):
        connection.execute("PRAGMA foreign_keys=ON")

    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as session:
        yield session
    await engine.dispose()


async def user(db, username, role=UserRole.teacher):
    row = User(
        tenant_id=11,
        username=username,
        hashed_password="synthetic-unusable",
        role=role,
        is_active=True,
    )
    db.add(row)
    await db.flush()
    return row


async def root(db, owner, *, class_id=101):
    row = WeeklyMonthlyPlan(
        tenant_id=11,
        owner_user_id=owner.id,
        teacher_id=owner.id,
        class_id=class_id,
        plan_kind=PlanKind.WEEKLY_ACTIVITY.value,
        current_version=1,
        revision=1,
    )
    db.add(row)
    await db.flush()
    version = WeeklyMonthlyPlanVersion(
        tenant_id=11,
        plan_id=row.id,
        version=1,
        plan_kind=PlanKind.WEEKLY_ACTIVITY.value,
        status="draft",
        week_start=date(2026, 9, 7),
        week_end=date(2026, 9, 13),
        week_number=2,
        grade="中班",
        class_name="合成班",
        teacher_names_json='["教师甲"]',
        theme_name="秋叶",
        canonical_payload_sha256="a" * 64,
        created_by=owner.id,
    )
    db.add(version)
    await db.flush()
    return row, version


def request(actor, plan, action=PlanAction.READ, **changes):
    values = {
        "action": action,
        "actor_id": actor.id,
        "actor_role": actor.role.value,
        "tenant_id": 11,
        "owner_teacher_id": plan.owner_user_id,
        "class_id": plan.class_id,
        "plan_kind": PlanKind.WEEKLY_ACTIVITY,
        "plan_id": plan.id,
        "plan_version": 1,
        "status": ReviewStatus.DRAFT,
    }
    return PlanAuthorizationRequest(**(values | changes))


async def daily(db, actor, text="旧晨谈"):
    return await save_daily_plan(
        db,
        11,
        actor.id,
        date(2026, 9, 7),
        2,
        "周一",
        "中班",
        "合成班",
        morning_talk_topic=text,
    )


async def test_gap_same_class_week_storage_allows_two_personal_roots(db):
    a, b = await user(db, "a"), await user(db, "b")
    await root(db, a)
    await root(db, b)
    ids = (await db.scalars(select(WeeklyMonthlyPlan.id))).all()
    assert len(ids) == 1, (
        "shared_unique_required: legacy stores two roots for one class/week"
    )


async def test_gap_owner_only_list_cannot_supply_self_and_other_candidates(db):
    a, b = await user(db, "a"), await user(db, "b")
    first, second = await daily(db, a), await daily(db, b)
    rows, total = await list_daily_plans_for_user(
        db,
        11,
        a.id,
        start_date=date(2026, 9, 7),
        end_date=date(2026, 9, 7),
    )

    assert (total, {row.id for row in rows}) == (2, {first.id, second.id}), (
        "class_candidate_service_required: owner list sees only one; "
        "do not broaden this legacy function or authorize by matching class_name"
    )


async def test_gap_split_discards_explicit_activity_name(monkeypatch):
    async def synthetic_ai(**_kwargs):
        return {
            "activity_name": "秋叶拼画",
            "activity_goal": "观察叶形",
            "activity_prep": "落叶",
            "activity_key": "观察",
            "activity_difficult": "表达",
            "activity_process": "观察落叶并拼画",
        }

    monkeypatch.setattr(lesson_plan_client, "call_ai", synthetic_ai)
    result = await lesson_plan_client.split_lesson_plan(
        raw_text="活动名称：秋叶拼画。观察落叶并拼画。",
        api_base_url="https://invalid.example",
        api_key="synthetic-unusable",
    )
    assert result.get("activity_name") == "秋叶拼画", (
        "split discards explicit source name"
    )


async def test_gap_activity_name_save_reload_and_revision(db):
    a = await user(db, "a")
    source = await save_daily_plan(
        db,
        11,
        a.id,
        date(2026, 9, 7),
        2,
        "周一",
        "中班",
        "合成班",
        activity_name="秋叶拼画",
    )
    await db.commit()
    loaded = await get_daily_plan_by_id_for_user(db, 11, a.id, source.id)
    assert loaded.activity_name == "秋叶拼画"
    assert loaded.revision == 1
    await save_daily_plan(
        db,
        11,
        a.id,
        source.plan_date,
        2,
        "周一",
        "中班",
        "合成班",
        expected_plan_id=source.id,
        expected_revision=1,
        activity_name="叶子拓印",
    )
    await db.commit()
    await db.refresh(loaded)
    assert (loaded.activity_name, loaded.revision) == ("叶子拓印", 2)


@pytest.mark.parametrize(
    "sunday,monday",
    [
        (date(2026, 9, 6), date(2026, 9, 7)),
        (date(2025, 12, 28), date(2025, 12, 29)),
    ],
)
def test_gap_preceding_sunday_needs_next_monday_identity(sunday, monday):
    semester_start = date(sunday.year, 9, 2)
    assert get_week_number(semester_start, sunday) == get_week_number(
        semester_start, monday
    ), (
        "teaching_anchor_required: natural-week function assigns Sunday to preceding week"
    )


@pytest.mark.parametrize(
    "day,label", [(date(2026, 9, 6), "周日"), (date(2026, 9, 12), "周六")]
)
def test_gap_legacy_day_rejects_supported_sixth_date(day, label):
    value = WeeklyDay(day, day.weekday(), label, "谈秋叶", "《秋叶》", "", "")
    assert value.day_date == day


@pytest.mark.parametrize(
    "field", ["weekly_focus", "environment_creation", "life_habits"]
)
def test_gap_two_items_are_not_rejected_by_legacy_narrative_contract(field):
    with pytest.raises(ValueError):
        replace(_canonical_weekly_plan(), **{field: "1.观察秋叶\n2.分享发现"})


async def test_gap_daily_game_renderer_repeats_weekly_fixed_game_count():
    plan = _canonical_weekly_plan()
    outdoor = "体能大循环\n集体游戏：1.《传球》目标①合作②轮流③安全\n2.《跳圈》目标①平衡②协调③坚持\n自主游戏：《沙水》目标①探索②整理③分享"
    plan = replace(
        plan, days=tuple(replace(day, outdoor_game=outdoor) for day in plan.days)
    )
    port = build_released_weekly_monthly_word_port()
    binding = await port.resolve_active(plan.scope.tenant_id, "weekly_activity_plan")
    rendered = await port.render(binding, plan)
    text = Document(BytesIO(rendered.rendered_bytes)).tables[0].cell(3, 2).text
    assert text.count("《传球》") == 1, (
        "weekly_games_required: old renderer repeats all five daily groups"
    )


async def test_baseline_cross_class_and_tenant_request_denied(db):
    a = await user(db, "a")
    plan, _ = await root(db, a)
    policy = DatabasePlanAuthorizationAdapter(db)
    assert (await policy.authorize(request(a, plan))).allowed
    for changes in ({"class_id": 102}, {"tenant_id": 12}):
        assert not (await policy.authorize(request(a, plan, **changes))).allowed


async def test_baseline_revoked_legacy_grant_is_rechecked(db):
    a = await user(db, "a")
    b = await user(db, "b", UserRole.teaching_admin)
    plan, version = await root(db, a)
    # Metadata fixture is not migration/immutable-trigger evidence.
    version.status = "submitted"
    await db.flush()
    repo = SqlAlchemyPlanAggregateReadRepository(db)
    scope = {
        "tenant_id": 11,
        "grantee_user_id": b.id,
        "teacher_user_id": a.id,
        "class_id": 101,
        "action": "read",
    }
    await repo.save_scope_grant(**scope, revision=1)
    policy = DatabasePlanAuthorizationAdapter(db)
    req = request(b, plan, status=ReviewStatus.SUBMITTED)
    assert (await policy.authorize(req)).allowed
    await repo.save_scope_grant(
        **scope,
        revision=2,
        expected_revision=1,
        is_active=False,
        revoked_at=datetime.now(UTC),
    )
    assert not (await policy.authorize(req)).allowed


async def test_baseline_source_change_does_not_rewrite_detached_weekly_body(db):
    a = await user(db, "a")
    source = await daily(db, a)
    plan = _canonical_weekly_plan()
    old = replace(
        plan,
        days=(
            replace(plan.days[0], morning_talk=source.morning_talk_topic),
            *plan.days[1:],
        ),
    )
    await save_daily_plan(
        db,
        11,
        a.id,
        source.plan_date,
        2,
        "周一",
        "中班",
        "合成班",
        expected_plan_id=source.id,
        expected_revision=source.revision,
        morning_talk_topic="新晨谈",
    )
    assert source.revision == 2
    assert old.days[0].morning_talk == "旧晨谈"


async def test_baseline_source_daily_write_does_not_cross_user(db):
    a, b = await user(db, "a"), await user(db, "b")
    source = await daily(db, a)
    assert await get_daily_plan_by_id_for_user(db, 11, b.id, source.id) is None
    assert source.morning_talk_topic == "旧晨谈"
