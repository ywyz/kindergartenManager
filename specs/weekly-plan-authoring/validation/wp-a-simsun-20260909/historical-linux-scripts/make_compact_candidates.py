from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt


OUT = Path("/home/ywyz/code/km-wpa-fonts-20260909/layout")


def _set_run_font(run, size: float = 12.0) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(key), "宋体")


def _format_paragraph(paragraph, *, title: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for run in paragraph.runs:
        _set_run_font(run, 16.0 if title else 12.0)
        if title:
            run.bold = True


def _set_cell(cell, text: str) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        _format_paragraph(paragraph)


def _build(kind: str) -> Path:
    source_name = {
        "normal": "normal-five-short.docx",
        "sunday": "preceding-sunday-six-short.docx",
        "saturday": "saturday-six-short.docx",
    }[kind]
    target = OUT / f"{kind}-compact-candidate.docx"
    document = Document(str(OUT / source_name))
    table = document.tables[0]
    day_count = 5 if kind == "normal" else 6
    for col in range(2, 2 + day_count):
        _set_cell(table.cell(1, col), "观察树叶并说颜色和形状。")
        _set_cell(table.cell(2, col), "春天颜色观察与同伴分享。")
    _set_cell(
        table.cell(3, 2),
        "体能大循环\n"
        "集体游戏：1.《接力跑》（目标：①奔跑。②合作。③守规。）\n"
        "2.《投沙包》（目标：①投掷。②瞄准。③轮流。）\n"
        "自主游戏：《小车道》（目标：①规划。②等待。③分享。）",
    )
    _set_cell(
        table.cell(4, 2),
        "1.户外游戏 2.区域游戏 3.专用室\n"
        "本周重点指导区域：建构区\n"
        "目标：\n1.搭建并说明。\n2.协商合作。\n3.记录分享。\n"
        "材料：积木、纸筒、图卡。\n"
        "指导：\n1.示范提问。\n2.鼓励调整。\n3.交流整理。",
    )
    _set_cell(table.cell(5, 2), "1.倾听并表达观察。\n2.协商轮流合作。\n3.记录发现过程。")
    _set_cell(table.cell(6, 2), "1.主题墙更新。\n2.图书角添材料。\n3.建构区留空间。")
    _set_cell(table.cell(7, 2), "1.洗手按步骤。\n2.进餐细嚼慢咽。\n3.整理材料归位。")
    _set_cell(table.cell(8, 2), "请家长和幼儿散步观察植物，说说颜色、形状和变化并带图分享。")
    for index, paragraph in enumerate(document.paragraphs[:3]):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_paragraph(paragraph, title=index == 0)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                _format_paragraph(paragraph)
    document.save(str(target))
    return target


if __name__ == "__main__":
    for kind in ("normal", "sunday", "saturday"):
        print(_build(kind))
