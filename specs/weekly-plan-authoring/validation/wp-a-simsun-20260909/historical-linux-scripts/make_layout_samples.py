from __future__ import annotations

import copy
import json
import shutil
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPO = Path("/home/ywyz/code/km-wpa-20260908")
TEMPLATE = REPO / "templates" / "weekplan.docx"
OUT = Path("/home/ywyz/code/km-wpa-fonts-20260909/layout")
USABLE_A4_WIDTH_TWIPS = 10466
FIVE_DAY_WIDTHS = [933, 920, 1722, 1722, 1723, 1723, 1723]
SIX_DAY_WIDTHS = [933, 920, 1435, 1435, 1435, 1436, 1436, 1436]


def _remove_after_first_table(document: Document) -> None:
    body = document._element.body
    first_table_seen = False
    for child in list(body):
        if child.tag == qn("w:tbl") and not first_table_seen:
            first_table_seen = True
            continue
        if first_table_seen and child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_run_font(run, size: float = 12.0, bold: bool | None = None) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(key), "宋体")


def _format_paragraph(paragraph, *, align=None, title=False) -> None:
    if align is not None:
        paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(20)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for run in paragraph.runs:
        _set_run_font(run, 16.0 if title else 12.0, True if title else None)


def _set_text(cell, text: str, *, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        _format_paragraph(paragraph, align=align)


def _set_col_widths(table, widths_twips: list[int]) -> None:
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if len(grid_cols) != len(widths_twips):
        raise RuntimeError(f"unexpected grid columns: {len(grid_cols)}")
    for grid_col, width in zip(grid_cols, widths_twips):
        grid_col.set(qn("w:w"), str(width))
    for tr in table._tbl.findall(qn("w:tr")):
        col = 0
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is None:
                continue
            span_node = tc_pr.find(qn("w:gridSpan"))
            span = int(span_node.get(qn("w:val"))) if span_node is not None else 1
            end = min(col + span, len(widths_twips))
            tcw = tc_pr.find(qn("w:tcW"))
            if tcw is None:
                from docx.oxml import OxmlElement

                tcw = OxmlElement("w:tcW")
                tc_pr.insert(0, tcw)
            tcw.set(qn("w:w"), str(sum(widths_twips[col:end])))
            tcw.set(qn("w:type"), "dxa")
            col = end


def _set_table_width(table, width_twips: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        from docx.oxml import OxmlElement

        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def _copy_new_column_borders(table) -> None:
    """Give add_column's new cells the retained rightmost-cell borders."""

    for tr in table._tbl.findall(qn("w:tr")):
        cells = tr.findall(qn("w:tc"))
        if len(cells) < 2:
            continue
        source_pr = cells[-2].find(qn("w:tcPr"))
        target_pr = cells[-1].find(qn("w:tcPr"))
        if source_pr is None or target_pr is None:
            continue
        source_borders = source_pr.find(qn("w:tcBorders"))
        if source_borders is None:
            continue
        target_borders = target_pr.find(qn("w:tcBorders"))
        if target_borders is not None:
            target_pr.remove(target_borders)
        target_pr.append(copy.deepcopy(source_borders))


def _short_content(kind: str) -> dict[str, str | list[str]]:
    if kind == "normal":
        labels = ["周一", "周二", "周三", "周四", "周五"]
    elif kind == "sunday":
        labels = ["周日", "周一", "周二", "周三", "周四", "周五"]
    elif kind == "saturday":
        labels = ["周一", "周二", "周三", "周四", "周五", "周六"]
    else:
        raise ValueError(kind)
    morning = [f"认识{chr(0x4E00 + i)}" for i in range(len(labels))]
    collective = [f"《活动{i + 1}》" for i in range(len(labels))]
    return {
        "morning": "\n".join(morning),
        "collective": "\n".join(collective),
        "outdoor": (
            "体能大循环\n"
            "集体游戏：1.《接力跑》（目标：①奔跑。②合作。③守规。）\n"
            "2.《投沙包》（目标：①投掷。②瞄准。③轮流。）\n"
            "自主游戏：《小车道》（目标：①规划。②等待。③分享。）"
        ),
        "area": (
            "1.户外游戏 2.区域游戏 3.专用室\n"
            "本周重点指导区域：建构区\n"
            "目标：\n1.搭建。\n2.表达。\n3.合作。\n"
            "材料：积木、纸筒。\n"
            "指导：\n1.示范。\n2.提问。\n3.分享。"
        ),
        "focus": "1.倾听。\n2.合作。\n3.表达。",
        "environment": "1.主题墙。\n2.图书角。\n3.材料架。",
        "life": "1.洗手：按步骤。\n2.进餐：细嚼。\n3.整理：归位。",
        "home": "请家长和幼儿一起观察春天并分享发现。",
        "labels": labels,
    }


def _long_content(kind: str) -> dict[str, str | list[str]]:
    if kind == "normal":
        labels = ["周一", "周二", "周三", "周四", "周五"]
    elif kind == "sunday":
        labels = ["周日", "周一", "周二", "周三", "周四", "周五"]
    elif kind == "saturday":
        labels = ["周一", "周二", "周三", "周四", "周五", "周六"]
    else:
        raise ValueError(kind)
    morning = [
        "围绕春天的树叶、花朵和天气变化交流观察方法，鼓励幼儿完整表达。"
        for _ in labels
    ]
    collective = [
        "《春天的颜色》：认识自然中的色彩变化，尝试用语言和图画记录发现，学习与同伴轮流分享。"
        for _ in labels
    ]
    return {
        "morning": "\n".join(morning),
        "collective": "\n".join(collective),
        "outdoor": (
            "体能大循环\n"
            "集体游戏：1.《森林接力探险》（目标：①练习起跑、绕障碍和安全停步，感受连续运动的节奏。"
            "②在小组合作中商量路线、轮流等待并为同伴加油。③尝试根据场地变化调整速度，体验完成任务的成就感。）\n"
            "2.《彩虹投掷站》（目标：①练习肩上投掷与双手接物，发展手眼协调。②观察距离和方向，学习用简单方法修正动作。"
            "③遵守轮流与安全规则，能够清楚表达自己的运动感受。）\n"
            "自主游戏：《小小城市建造师》（目标：①共同设计道路、桥梁和休息站，发展空间想象。②遇到连接困难时尝试协商、分工与调整。"
            "③愿意介绍作品并倾听同伴建议，在持续游戏中保持专注。）"
        ),
        "area": (
            "1.户外游戏 2.区域游戏 3.专用室\n"
            "本周重点指导区域：建构区\n"
            "目标：\n"
            "1.根据生活经验设计有入口、通道和功能区域的立体建筑，能够说明自己的构思。\n"
            "2.与同伴协商材料分配和空间使用，在结构不稳时共同寻找原因并调整方案。\n"
            "3.用图示、符号或简短语言记录搭建过程，愿意比较不同方案并尊重同伴作品。\n"
            "材料：木质积木、连接件、纸筒、彩纸、道路标志卡、记录纸和水彩笔。\n"
            "指导：\n"
            "1.先观察幼儿的设计意图，再用开放问题支持其解释结构、材料和功能之间的关系。\n"
            "2.提醒幼儿测试承重与通行效果，鼓励小组在保留原想法的基础上提出至少一种改进办法。\n"
            "3.游戏结束前组织同伴参观和简短分享，引导幼儿整理材料并留下可继续发展的记录。"
        ),
        "focus": (
            "1.持续关注幼儿在观察记录中的细节表达，支持其使用完整语句描述季节变化。\n"
            "2.在合作游戏中引导幼儿协商角色、轮流使用材料，并学习用语言解决小分歧。\n"
            "3.鼓励幼儿根据自己的问题选择材料、验证想法，再用图画或符号留下过程证据。"
        ),
        "environment": (
            "1.在主题墙设置叶片、花朵、天气和幼儿观察记录的可替换展示区，保留日期线索。\n"
            "2.在图书角补充春季自然观察、合作故事和图画记录材料，设置安静交流的小桌面。\n"
            "3.在建构区按形状与功能分类摆放积木、连接件和标志卡，留出作品继续搭建的空间。"
        ),
        "life": (
            "1.洗手：按照步骤打湿、起泡、揉搓、冲洗和擦干，能够在提醒后检查指缝是否洗净。\n"
            "2.进餐：坐姿稳定、细嚼慢咽，尝试自主收拾餐具并用合适音量表达需要。\n"
            "3.整理：游戏结束前先分类再归位，学习检查公共材料是否完整并与同伴共同完成整理。"
        ),
        "home": (
            "请家长利用散步或周末时间和幼儿观察一种春天的植物，鼓励幼儿说出颜色、形状和变化，"
            "并把一段简短描述或一张观察图带到班级分享。"
        ),
        "labels": labels,
    }


def _fill_sample(kind: str, long: bool, out_path: Path) -> None:
    document = Document(str(TEMPLATE))
    _remove_after_first_table(document)
    table = document.tables[0]
    content = _long_content(kind) if long else _short_content(kind)
    labels = content["labels"]
    assert isinstance(labels, list)
    start = date(2026, 9, 7)
    if kind == "sunday":
        dates = [start - timedelta(days=1), *[start + timedelta(days=i) for i in range(5)]]
    elif kind == "saturday":
        dates = [start + timedelta(days=i) for i in range(6)]
    else:
        dates = [start + timedelta(days=i) for i in range(5)]
    display_start, display_end = dates[0], dates[-1]

    if len(labels) == 6:
        table.add_column(Inches(1.0))
        _copy_new_column_borders(table)
        widths = SIX_DAY_WIDTHS
        for row in (3, 4, 5, 6, 7, 8):
            table.cell(row, 2).merge(table.cell(row, 7))
    else:
        widths = FIVE_DAY_WIDTHS
    _set_col_widths(table, widths)
    _set_table_width(table, sum(widths))
    if sum(widths) > USABLE_A4_WIDTH_TWIPS:
        raise AssertionError(sum(widths))

    theme = "《春天里的发现》" if not long else "《春天里的发现与城市建造》"
    header = (
        f"主题名称：{theme}    班级：中一班    第（2）周（"
        f"{display_start:%Y年%m月%d日}——{display_end:%Y年%m月%d日}）"
    )
    people = "教师：张老师、李老师、王老师    保育员：赵阿姨"
    document.paragraphs[0].text = "幼儿园每周工作计划表"
    document.paragraphs[1].text = header
    document.paragraphs[2].text = people
    for i, paragraph in enumerate(document.paragraphs[:3]):
        _format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, title=i == 0)

    _set_text(table.cell(0, 0), "第2周", align=WD_ALIGN_PARAGRAPH.CENTER)
    for col, (label, day_date) in enumerate(zip(labels, dates), start=2):
        _set_text(table.cell(0, col), f"{label}\n{day_date:%m月%d日}", align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_text(table.cell(1, col), content["morning"].split("\n")[col - 2])
        _set_text(table.cell(2, col), content["collective"].split("\n")[col - 2])
    _set_text(table.cell(3, 2), content["outdoor"])
    _set_text(table.cell(4, 2), content["area"])
    for row, key in ((5, "focus"), (6, "environment"), (7, "life"), (8, "home")):
        _set_text(table.cell(row, 2), content[key])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                _format_paragraph(paragraph, align=paragraph.alignment)
                for run in paragraph.runs:
                    if run.text and not run.font.size:
                        _set_run_font(run)
    document.save(str(out_path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    reference = OUT / "weekplan.docx"
    shutil.copy2(TEMPLATE, reference)
    specs = [
        ("normal", False, "normal-five-short.docx"),
        ("normal", True, "normal-five-long.docx"),
        ("sunday", False, "preceding-sunday-six-short.docx"),
        ("sunday", True, "preceding-sunday-six-long.docx"),
        ("saturday", False, "saturday-six-short.docx"),
        ("saturday", True, "saturday-six-long.docx"),
    ]
    for kind, long, filename in specs:
        _fill_sample(kind, long, OUT / filename)
    print(json.dumps({"template": str(TEMPLATE), "outputs": [x[2] for x in specs]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
