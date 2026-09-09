"""WP-B synthetic daily-name regression; no network or business database."""

import json
from datetime import date
from io import BytesIO

import pytest
from docx import Document
from sqlalchemy.orm.exc import StaleDataError

from app.api.schemas import DailyPlanOut
from app.core.exceptions import AiParseError
from app.integration.ai_client import lesson_plan_client as client
from app.integration.word_export.exporter import export_daily_plan
from app.repository.daily_plan_repository import (
    get_daily_plan_by_id_for_user,
    save_daily_plan,
)
from app.service.agent.confirmed_write import _daily_plan_snapshot

FIELDS = {
    "activity_goal": "目标",
    "activity_prep": "准备",
    "activity_key": "重点",
    "activity_difficult": "难点",
    "activity_process": "过程",
}


async def split(monkeypatch, response, raw="活动名称：秋叶拼画", prompt=None):
    async def call(**kwargs):
        if prompt is not None:
            assert kwargs["messages"][0]["content"] == prompt
        return response.copy()

    monkeypatch.setattr(client, "call_ai", call)
    return await client.split_lesson_plan(
        raw, "https://invalid.example", "synthetic", system_prompt=prompt
    )


async def test_explicit_name_preserved(monkeypatch):
    assert (await split(monkeypatch, FIELDS | {"activity_name": "秋叶拼画"}))[
        "activity_name"
    ] == "秋叶拼画"


@pytest.mark.parametrize(
    "raw", ["活动目标：观察叶形\n活动过程：秋叶拼画\n反思：开心", "活动目标：秋叶拼画"]
)
async def test_no_title_does_not_invent_from_body(monkeypatch, raw):
    result = await split(monkeypatch, FIELDS | {"activity_name": "秋叶拼画"}, raw)
    assert result["activity_name"] == ""


async def test_legacy_custom_prompt_missing_name_is_empty(monkeypatch):
    result = await split(monkeypatch, FIELDS, prompt="用户原版五字段提示词")
    assert result["activity_name"] == ""


async def test_default_missing_name_rejected(monkeypatch):
    with pytest.raises(AiParseError):
        await split(monkeypatch, FIELDS)


@pytest.mark.parametrize("value", [None, 7, [], {}, True, "叶" * 86])
async def test_invalid_name_type_or_utf8_limit_rejected(monkeypatch, value):
    with pytest.raises(AiParseError):
        await split(monkeypatch, FIELDS | {"activity_name": value})


async def save(session, **kwargs):
    return await save_daily_plan(
        session, 11, 7, date(2026, 9, 7), 2, "周一", "中班", "合成班", **kwargs
    )


async def test_name_save_reload_api_word_and_snapshot(async_session):
    plan = await save(async_session, activity_name="秋叶拼画", activity_goal="目标保持")
    await async_session.commit()
    plan_id = plan.id
    await save(
        async_session,
        expected_plan_id=plan_id,
        expected_revision=1,
        activity_name="叶子拓印",
    )
    await async_session.commit()
    async_session.expunge_all()
    plan = await get_daily_plan_by_id_for_user(async_session, 11, 7, plan_id)
    assert (plan.activity_name, plan.revision, plan.activity_goal) == (
        "叶子拓印",
        2,
        "目标保持",
    )
    assert DailyPlanOut.from_model(plan).model_dump()["activity_name"] == "叶子拓印"
    doc = Document(BytesIO(export_daily_plan(plan, [])))
    assert doc.tables[0].rows[6].cells[-1].text == "活动主题：叶子拓印"
    assert "目标保持" in doc.tables[0].rows[7].cells[-1].text
    snapshot, _digest = _daily_plan_snapshot(plan)
    assert json.loads(snapshot)["snapshot_schema_version"] == 2
    assert json.loads(snapshot)["activity_name"] == "叶子拓印"
    await save(
        async_session,
        expected_plan_id=plan_id,
        expected_revision=2,
        activity_name="叶子拓印",
    )
    assert plan.revision == 2
    with pytest.raises(StaleDataError):
        await save(
            async_session,
            expected_plan_id=plan_id,
            expected_revision=1,
            activity_name="旧页面",
        )
    for tenant, user in [(12, 7), (11, 8)]:
        assert (
            await get_daily_plan_by_id_for_user(async_session, tenant, user, plan_id)
            is None
        )
        with pytest.raises(StaleDataError):
            await save_daily_plan(
                async_session,
                tenant,
                user,
                date(2026, 9, 7),
                2,
                "周一",
                "中班",
                "合成班",
                expected_plan_id=plan_id,
                expected_revision=2,
                activity_name="越权",
            )


async def test_null_old_row_readable(async_session):
    plan = await save(async_session)
    assert DailyPlanOut.from_model(plan).model_dump()["activity_name"] is None
    doc = Document(BytesIO(export_daily_plan(plan, [])))
    assert doc.tables[0].rows[6].cells[-1].text == "活动主题："


@pytest.mark.parametrize("value", [1, True, [], "叶" * 86])
async def test_repository_rejects_bad_name(async_session, value):
    with pytest.raises(ValueError):
        await save(async_session, activity_name=value)


@pytest.mark.parametrize(
    "raw,name",
    [
        ("活动名称：《秋叶拼画》", "《秋叶拼画》"),
        ("集体活动名称：秋叶拼画", "秋叶拼画"),
        ("《秋叶拼画》\n活动目标：观察", "秋叶拼画"),
        ("秋叶拼画\n活动目标：观察", "秋叶拼画"),
    ],
)
async def test_explicit_title_formats(monkeypatch, raw, name):
    assert (await split(monkeypatch, FIELDS | {"activity_name": name}, raw))[
        "activity_name"
    ] == name


@pytest.mark.parametrize(
    "raw",
    [
        "活动目标：观察叶形",
        "活动过程：观察落叶",
        "活动反思：孩子开心",
        "观察落叶并拼画。",
    ],
)
async def test_entire_body_line_cannot_be_title(monkeypatch, raw):
    assert (await split(monkeypatch, FIELDS | {"activity_name": raw}, raw))[
        "activity_name"
    ] == ""


async def test_name_word_preserves_whitespace_exactly(async_session):
    name = "  秋叶拼画  \n第二行 "
    plan = await save(async_session, activity_name=name)
    await async_session.commit()
    assert DailyPlanOut.from_model(plan).activity_name == name
    text = (
        Document(BytesIO(export_daily_plan(plan, []))).tables[0].rows[6].cells[-1].text
    )
    assert text == "活动主题：" + name


async def test_invalid_unicode_is_sanitized_parse_error(monkeypatch):
    with pytest.raises(AiParseError, match="256 UTF-8"):
        await split(monkeypatch, FIELDS | {"activity_name": "\ud800"})


async def test_blank_source_title_stays_empty(monkeypatch):
    result = await split(
        monkeypatch, FIELDS | {"activity_name": "   "}, "活动名称：\n活动目标：观察"
    )
    assert result["activity_name"] == ""
