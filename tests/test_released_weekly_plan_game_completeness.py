"""Regression for content lost while filling the released merged game rows.

These renderer tests are not protected UI or Office acceptance evidence.
"""

from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document

from app.integration.word_export.released_weekly_monthly_word_port import (
    build_released_weekly_monthly_word_port,
)
from app.service.weekly_monthly_plans.qualification_orchestration import (
    _canonical_weekly_plan,
)


@pytest.mark.asyncio
@pytest.mark.parametrize("line_count", [1, 30], ids=["normal", "long_chinese"])
async def test_released_weekly_game_rows_preserve_every_day(line_count: int) -> None:
    plan = _canonical_weekly_plan()
    days = tuple(
        replace(
            day,
            outdoor_game="\n".join(
                f"{day.weekday_cn}户外游戏第{index + 1}段：观察、合作与表达。"
                for index in range(line_count)
            ),
            area_game="\n".join(
                f"{day.weekday_cn}区域游戏第{index + 1}段：比较、记录与分享。"
                for index in range(line_count)
            ),
        )
        for day in plan.days
    )
    plan = replace(plan, days=days)
    port = build_released_weekly_monthly_word_port()
    binding = await port.resolve_active(plan.scope.tenant_id, "weekly_activity_plan")

    rendered = await port.render(binding, plan)

    document = Document(BytesIO(rendered.rendered_bytes))
    for table in document.tables:
        for row, field in ((3, "outdoor_game"), (4, "area_game")):
            cell = table.cell(row, 2)
            assert all(
                table.cell(row, column)._tc is cell._tc for column in range(2, 7)
            )
            positions = []
            for day in days:
                value = getattr(day, field)
                assert value in cell.text
                assert cell.text.count(value) == 1
                positions.append(cell.text.index(value))
            assert positions == sorted(positions)


@pytest.mark.asyncio
async def test_released_weekly_game_rows_keep_empty_days_distinct() -> None:
    plan = _canonical_weekly_plan()
    days = tuple(
        replace(
            day,
            outdoor_game="" if index == 2 else f"{day.weekday_cn}观察秋叶",
            area_game="" if index == 2 else f"{day.weekday_cn}合作搭建",
        )
        for index, day in enumerate(plan.days)
    )
    plan = replace(plan, days=days)
    port = build_released_weekly_monthly_word_port()
    binding = await port.resolve_active(plan.scope.tenant_id, "weekly_activity_plan")

    rendered = await port.render(binding, plan)

    for table in Document(BytesIO(rendered.rendered_bytes)).tables:
        for row, field in ((3, "outdoor_game"), (4, "area_game")):
            lines = table.cell(row, 2).text.splitlines()
            assert len(lines) == 5
            assert lines[2] == "周三："
            assert "None" not in table.cell(row, 2).text
            for index, day in enumerate(days):
                assert lines[index] == f"{day.weekday_cn}：{getattr(day, field)}"
