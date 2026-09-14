"""Personal source isolation and non-destructive weekly filling."""

from datetime import date
from uuid import uuid4

import pytest
from sqlalchemy import insert, update

from app.integration.ai_client import weekly_authoring_client as client
from app.repository.shared_weekly_repository import VERSION
from app.repository.source_mapping_repository import DAILY, MAPPING
from app.service.shared_weekly.authoring_application import (
    AuthoringApplication,
    SlotChange,
)
from app.service.shared_weekly.autofill import morning_summary
from app.service.shared_weekly.root_contracts import WeeklyThemeDraft
from tests.test_wpc_identity import world as _world
from tests.test_wpc_shared_root import ready, rows

world = _world


async def setup(world):
    apps, scope, _ = await ready(world)
    result = await apps[3].create(world[2][3], scope, WeeklyThemeDraft("秋天"), uuid4())
    app = AuthoringApplication(world[0], lambda: world[1][3])
    edit = await app.begin_authoring(world[2][3], result.plan.plan_id)
    return app, edit


async def daily(
    world,
    *,
    owner=3,
    class_name="四班",
    day=date(2026, 9, 9),
    name="落叶",
    topic="秋天的变化。你看到了什么？",
    morning_activity="",
):
    async with world[0]() as session:
        row = await session.execute(
            insert(DAILY).values(
                tenant_id=11,
                user_id=owner,
                plan_date=day,
                week_number=2,
                weekday_cn="周三",
                grade="中班",
                class_name=class_name,
                activity_name=name,
                morning_talk_topic=topic,
                morning_talk_questions="为什么叶子会变黄？",
                morning_activity=morning_activity,
                outdoor_activity="集体游戏：《跳圈》（目标：①双脚跳跃②保持平衡③遵守规则）\n自主游戏：《玩沙》（目标：①探索沙土②使用工具③合作游戏）",
                indoor_area="本周重点指导区域：建构区\n活动目标：\n1.合作搭建\n2.认识形状\n3.表达想法\n材料：木积木与纸盒\n指导要点：\n1.观察幼儿\n2.提供支持\n3.鼓励交流",
            )
        )
        await session.commit()
        return row.inserted_primary_key[0]


async def test_unmapped_personal_autofill_correct_cells_no_writes(world):
    app, edit = await setup(world)
    sid = await daily(world)
    await daily(world, owner=4, name="其他老师")
    await daily(world, class_name="其他班", name="其他班活动")
    await daily(world, day=date(2026, 9, 12), name="周外活动")
    before, versions = await rows(world, DAILY), await rows(world, VERSION)
    filled = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    d = next(d for d in filled.body.days if d.day == date(2026, 9, 9))
    assert (d.activity_name, d.morning_talk_topic, d.morning_talk_questions) == (
        "落叶",
        "秋天的变化。",
        "",
    )
    assert d.outdoor_activity == d.indoor_area == ""
    assert filled.body.value_at("games.collective.0.name") == "跳圈"
    assert filled.body.value_at("area.materials") == "木积木与纸盒"
    assert filled.body.value_at("area.guidance.2") == "鼓励交流"
    assert {s.source_id for s in filled.body.sources} == {sid}
    assert await rows(world, DAILY) == before
    assert await rows(world, VERSION) == versions
    assert await rows(world, MAPPING) == ()
    saved = await app.save_edit(world[2][3], filled.page_id, filled.page, uuid4())
    other = AuthoringApplication(world[0], lambda: world[1][4])
    shared = await other.begin_authoring(world[2][4], saved.plan_id)
    assert shared.body.days == filled.body.days
    assert [s.value for s in shared.body.slots] == [s.value for s in filled.body.slots]
    async with other.source_transaction(world[2][4], shared.target) as context:
        visible, _ = await other._visible(context)
        assert sid not in {s.source_id for s in visible}


async def test_manual_and_saved_slots_preserved_on_reopen_and_changed_source(world):
    app, edit = await setup(world)
    sid = await daily(world)
    edit = await app.update_slots(
        world[2][3],
        edit.page_id,
        edit.page,
        (SlotChange("area.materials", "手填材料"),),
    )
    edit = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    saved = await app.save_edit(world[2][3], edit.page_id, edit.page, uuid4())
    async with world[0]() as session:
        await session.execute(
            update(DAILY)
            .where(DAILY.c.id == sid)
            .values(activity_name="变更后", revision=2)
        )
        await session.commit()
    reopened = await app.begin_authoring(world[2][3], saved.plan_id)
    filled = await app.autofill_owned(world[2][3], reopened.page_id, reopened.page)
    assert filled.body == reopened.body
    assert filled.body.value_at("area.materials") == "手填材料"


async def test_conflicting_same_day_requires_choice_not_highest_revision(world):
    app, edit = await setup(world)
    await daily(world, name="甲")
    sid = await daily(world, name="乙")
    filled = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    assert not any(d.activity_name for d in filled.body.days)
    assert len(app.owned_conflicts(world[2][3], edit.page_id)) == 1
    chosen = await app.autofill_owned(world[2][3], filled.page_id, filled.page, (sid,))
    assert (
        next(d for d in chosen.body.days if d.day == date(2026, 9, 9)).activity_name
        == "乙"
    )


@pytest.mark.parametrize(
    "text,expected",
    [
        ("你喜欢什么", ""),
        ("为什么树叶变黄", ""),
        ("秋天。问题：你看到了什么？", "秋天。"),
        ("主题：秋天的颜色", "秋天的颜色"),
    ],
)
def test_morning_topics_do_not_retain_questions_without_punctuation(text, expected):
    assert morning_summary(text) == expected


async def test_active_managed_prompt_reaches_ai_and_question_output_rejected(
    world, monkeypatch
):
    from app.repository.prompt_repository import save_new_version

    app, edit = await setup(world)
    await daily(world)
    edit = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    async with world[0]() as session:
        await save_new_version(session, 11, 3, "weekly_focus", "本次激活提示词")
    captured = []

    async def config(*args):
        return client.WeeklyAIConfig(
            "https://synthetic.invalid", "synthetic", "synthetic"
        )

    async def generate(prompt, payload, configuration):
        captured.append(prompt)
        return {"values": {p: "内容" + p for p in payload["fields"]}}

    monkeypatch.setattr(client, "load_config", config)
    monkeypatch.setattr(client, "generate", generate)
    await app.generate_missing(world[2][3], edit.page_id, edit.page, "weekly_focus")
    assert captured == ["本次激活提示词"]


async def test_personal_snapshot_later_explicit_mapping_is_not_broken(world):
    from app.service.shared_weekly.mapping_application import SourceMappingApplication
    from app.service.shared_weekly.mapping_contracts import MappingTarget

    app, edit = await setup(world)
    sid = await daily(world)
    edit = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    await app.save_edit(world[2][3], edit.page_id, edit.page, uuid4())
    mapping = SourceMappingApplication(world[0], lambda: world[1][2])
    scope = edit.target.authorization.scope
    preview = await mapping.preview(
        world[2][2], MappingTarget(sid, scope.class_instance_id, scope.semester_id)
    )
    result = await mapping.confirm(
        world[2][2], preview.candidate_id, confirmed=True, operation_id=uuid4()
    )
    assert result.revision == 2
    import asyncio

    from alembic.config import Config

    from alembic import command

    with pytest.raises(
        RuntimeError, match="personal_snapshot_downgrade_requires_verified_restore"
    ):
        await asyncio.to_thread(
            command.downgrade, Config("alembic.ini"), "d375e9ab2148"
        )


def test_export_only_topic_and_existing_materials(tmp_path):
    from io import BytesIO

    from docx import Document

    from app.integration.word_export.shared_weekly_word import SEED_PATH, fill_document
    from app.service.shared_weekly.layout_contracts import WeekDisplay
    from tests.test_wpe_word_layout import complete_body

    body = complete_body()
    data = fill_document(
        SEED_PATH.read_bytes(), body, WeekDisplay("小一班", "第一学期", 3)
    )
    document = Document(BytesIO(data))
    assert document.tables[0].cell(1, 2).text == body.days[0].morning_talk_topic
    assert body.days[0].morning_talk_questions not in document.tables[0].cell(1, 2).text
    assert body.value_at("area.materials") in document.tables[0].cell(4, 2).text
    (tmp_path / "weekly.docx").write_bytes(data)


async def test_ai_question_without_question_mark_rejected(world, monkeypatch):
    from app.integration.ai_client import weekly_authoring_client as client
    from app.service.academic_identity.contracts import IdentityRejected

    app, edit = await setup(world)
    await daily(world)
    edit = await app.autofill_owned(world[2][3], edit.page_id, edit.page)

    async def config(*args):
        return client.WeeklyAIConfig(
            "https://synthetic.invalid", "synthetic", "synthetic"
        )

    async def generate(prompt, payload, configuration):
        return {"values": {p: "你喜欢什么颜色" for p in payload["fields"]}}

    monkeypatch.setattr(client, "load_config", config)
    monkeypatch.setattr(client, "generate", generate)
    with pytest.raises(IdentityRejected, match="ai_invalid"):
        await app.regenerate(
            world[2][3],
            edit.page_id,
            edit.page,
            "weekly_morning_talk",
            ("days.2026-09-09.morning_talk_topic",),
        )
    assert app._page(world[2][3], edit.page_id).view == edit


async def test_cleared_imported_slot_stays_empty_after_save_reopen(world):
    app, edit = await setup(world)
    await daily(world)
    edit = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    edit = await app.update_slots(
        world[2][3], edit.page_id, edit.page, (SlotChange("area.materials", ""),)
    )
    saved = await app.save_edit(world[2][3], edit.page_id, edit.page, uuid4())
    reopened = await app.begin_authoring(world[2][3], saved.plan_id)
    filled = await app.autofill_owned(world[2][3], reopened.page_id, reopened.page)
    assert filled.body.value_at("area.materials") == ""


async def test_question_only_daily_variants_do_not_require_choice(world):
    app, edit = await setup(world)
    await daily(world, topic="秋天。问题：为什么落叶？")
    await daily(world, topic="秋天。问题：你看到什么？")
    filled = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    assert not app.owned_conflicts(world[2][3], edit.page_id)
    assert any(d.morning_talk_topic == "秋天。" for d in filled.body.days)


async def test_daily_outdoor_area_format_reuses_named_game_goals(world):
    from app.service.shared_weekly.source_structure import extract_options

    app, edit = await setup(world)
    sid = await daily(world)
    async with world[0]() as session:
        await session.execute(
            update(DAILY)
            .where(DAILY.c.id == sid)
            .values(
                revision=2,
                outdoor_activity="游戏区域：沙水区 、 攀爬区\n重点指导：攀爬区\n活动目标：\n1.保持平衡\n2.尝试攀爬\n3.遵守规则\n指导要点：\n1.观察\n2.支持\n3.鼓励",
            )
        )
        await session.commit()
    edit = await app.update_slots(
        world[2][3],
        edit.page_id,
        edit.page,
        (SlotChange("games.autonomous.name", "攀爬区"),),
    )
    filled = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    assert filled.body.value_at("games.autonomous.goals.0") == "保持平衡"
    options = extract_options(filled.body)
    assert any(o.kind == "outdoor" and o.name == "攀爬区" for o in options)


async def test_saved_manual_clear_without_source_reference_is_preserved(world):
    app, edit = await setup(world)
    edit = await app.update_slots(
        world[2][3],
        edit.page_id,
        edit.page,
        (SlotChange("area.materials", "手工材料"),),
    )
    edit = await app.update_slots(
        world[2][3], edit.page_id, edit.page, (SlotChange("area.materials", ""),)
    )
    saved = await app.save_edit(world[2][3], edit.page_id, edit.page, uuid4())
    await daily(world)
    reopened = await app.begin_authoring(world[2][3], saved.plan_id)
    filled = await app.autofill_owned(world[2][3], reopened.page_id, reopened.page)
    assert filled.body.value_at("area.materials") == ""


async def test_actual_daily_morning_activity_fills_games_and_saves_source(world):
    app, edit = await setup(world)
    sid = await daily(world)
    await daily(world, owner=4, morning_activity="集体游戏：其他教师游戏")
    await daily(world, class_name="其他班", morning_activity="集体游戏：其他班游戏")
    raw = "体能大循环：\n集体游戏：跳圈\n自主游戏：玩沙\n重点指导：玩沙\n活动目标：\n1.探索沙土\n2.使用工具\n3.合作游戏\n指导要点：\n1.观察\n2.支持\n3.鼓励"
    async with world[0]() as session:
        await session.execute(
            update(DAILY)
            .where(DAILY.c.id == sid)
            .values(
                revision=2,
                morning_activity=raw,
                outdoor_activity="游戏区域：沙水区 、 攀爬区\n重点指导：攀爬区\n活动目标：\n1.保持平衡\n2.尝试攀爬\n3.遵守规则",
            )
        )
        await session.commit()
    filled = await app.autofill_owned(world[2][3], edit.page_id, edit.page)
    assert filled.body.value_at("games.collective.0.name") == "跳圈"
    assert filled.body.value_at("games.autonomous.name") == "玩沙"
    assert filled.body.value_at("games.autonomous.goals.0") == "探索沙土"
    assert filled.body.value_at("games.collective.0.goals.0") == ""
    snapshot = next(
        s for s in filled.body.sources if s.source_field == "morning_activity"
    )
    assert snapshot.imported_value == raw
    again = await app.autofill_owned(world[2][3], filled.page_id, filled.page)
    assert again.body == filled.body
    saved = await app.save_edit(world[2][3], again.page_id, again.page, uuid4())
    reopened = await app.begin_authoring(world[2][3], saved.plan_id)
    assert [s.value for s in reopened.body.slots] == [
        s.value for s in filled.body.slots
    ]
    assert (
        next(
            s for s in reopened.body.sources if s.source_field == "morning_activity"
        ).imported_value
        == raw
    )
    assert (await rows(world, DAILY))[0]["morning_activity"] == raw


def test_daily_focus_goals_belong_to_named_collective_game():
    from app.service.shared_weekly.source_structure import _games

    result = _games(
        "体能大循环：\n集体游戏：跳圈\n自主游戏：玩沙\n重点指导：跳圈\n活动目标：\n1.双脚跳跃\n2.保持平衡\n3.遵守规则"
    )
    assert ("goals.0", "双脚跳跃") in result[0][2]
    assert result[1][2] == (("name", "玩沙"),)
