"""Initial coverage: new independent v3 layout, not historical business RED."""

import re
from dataclasses import replace
from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.integration.word_export.shared_weekly_word import (
    SEED_PATH,
    SEED_SHA256,
    LayoutRejected,
    SharedWeeklyWordPort,
    _rendered_grid,
    fill_document,
)
from app.service.shared_weekly.authoring_contracts import (
    AuthoringCalendar,
    AuthoringSlot,
    WeeklyAuthoringDraft,
)
from app.service.shared_weekly.body_contracts import (
    CollaborationDay,
    WeeklyCollaborationDraft,
)
from app.service.shared_weekly.layout_contracts import WeekDisplay
from app.service.shared_weekly.people_contracts import People


def complete_body(count=5, long=False):
    days = tuple(
        CollaborationDay(
            date(2026, 9, 14) + timedelta(days=i), "晨谈", "问题", "名称", "", ""
        )
        for i in range(count)
    )
    base = WeeklyCollaborationDraft(
        "秋天", People(("甲老师", "乙老师"), "丙老师"), days, ()
    )
    body = WeeklyAuthoringDraft.from_collaboration(base)
    body = body.with_slots(
        tuple(
            (path, AuthoringSlot(("内容" * 30 if long else "项") + str(i)))
            for i, path in enumerate(body.paths)
        )
    )
    return replace(
        body,
        calendar=AuthoringCalendar(
            "test.v1",
            "a" * 64,
            tuple(
                (d.day, i != 3, "" if i != 3 else "中秋节放假")
                for i, d in enumerate(days)
            ),
        ),
    )


@pytest.mark.parametrize("count", [5, 6])
def test_controlled_seed_fill_roundtrip_fixed_format(count):
    body = complete_body(count)
    seed = SEED_PATH.read_bytes()
    assert sha256(seed).hexdigest() == SEED_SHA256
    doc = Document(
        BytesIO(fill_document(seed, body, WeekDisplay("小一班", "第一学期", 3)))
    )
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.columns) == count + 2
    assert len(table.rows) == 9
    assert table.cell(1, 5).text == "中秋节放假"
    assert table.cell(2, 5).text == ""
    assert "《秋天》" in doc.paragraphs[1].text
    assert "甲老师、乙老师" in doc.paragraphs[2].text
    assert doc.paragraphs[0].runs[0].font.size.pt == 16
    paragraphs = list(doc.paragraphs[1:]) + [
        p for r in table.rows for c in r.cells for p in c.paragraphs
    ]
    for p in paragraphs:
        assert p.paragraph_format.line_spacing.pt == 20
        for run in p.runs:
            assert run.font.size.pt == 12
            assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Serif CJK SC"
    assert sha256(SEED_PATH.read_bytes()).hexdigest() == SEED_SHA256


@pytest.mark.parametrize("count", [5, 6])
@pytest.mark.real_render
async def test_real_local_render_long_fails_without_delivery(count):
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(count, long=True), WeekDisplay("小一班", "第一学期", 3)
    )
    assert result.reason == "layout_overflow"
    assert result.pages > 1
    assert not result.fits and result.data is None
    assert result.binding.active_version == "local-unqualified"


async def test_formal_qualification_not_inherited():
    with pytest.raises(ValueError, match="qualification_required"):
        await SharedWeeklyWordPort().resolve_binding(1)


async def test_incomplete_draft_refused_before_renderer():
    body = complete_body().with_slot("home", AuthoringSlot())
    result = await SharedWeeklyWordPort().inspect_local(
        body, WeekDisplay("小一班", "第一学期", 3)
    )
    assert not result.fits and result.pages == 0 and result.data is None


def test_consecutive_holiday_continuation_keeps_blank_cell():
    body = complete_body()
    columns = tuple(
        (day, False, "中秋节放假" if i == 3 else "")
        if i >= 3
        else (day, teaching, label)
        for i, (day, teaching, label) in enumerate(body.calendar.columns)
    )
    body = replace(body, calendar=replace(body.calendar, columns=columns))
    document = Document(
        BytesIO(
            fill_document(
                SEED_PATH.read_bytes(), body, WeekDisplay("小一班", "第一学期", 3)
            )
        )
    )
    assert document.tables[0].cell(1, 5).text == "中秋节放假"
    assert document.tables[0].cell(2, 5).text == ""
    assert document.tables[0].cell(1, 6).text == ""
    assert document.tables[0].cell(2, 6).text == ""


@pytest.mark.parametrize("count", [5, 6])
@pytest.mark.real_render
async def test_local_short_reports_actual_render(count):
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(count), WeekDisplay("小一班", "第一学期", 3)
    )
    print({"columns": count, "pages": result.pages, "reason": result.reason})
    assert result.pages >= 1
    assert result.reason in ("fits", "layout_overflow")
    assert (result.data is not None) == result.fits


@pytest.mark.real_render
async def test_six_short_complete_content_renders_on_one_page():
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(6), WeekDisplay("小一班", "第一学期", 3)
    )
    assert result.fits, (result.pages, result.reason)
    assert result.data is not None


def test_emitted_table_keeps_controlled_template_cell_borders():
    doc = Document(
        BytesIO(
            fill_document(
                SEED_PATH.read_bytes(),
                complete_body(6),
                WeekDisplay("小一班", "第一学期", 3),
            )
        )
    )
    for row in doc.tables[0].rows:
        for cell in row.cells:
            borders = cell._tc.tcPr.find(qn("w:tcBorders"))
            assert borders is not None, "visible fixed template grid missing"
            assert all(
                borders.find(qn("w:" + side)) is not None
                for side in ("top", "left", "bottom", "right")
            )


@pytest.mark.real_render
@pytest.mark.real_render
async def test_actual_render_missing_duplicate_label_is_not_delivered(monkeypatch):
    import xml.etree.ElementTree as ET

    from app.integration.word_export import shared_weekly_word as module

    original = module._process

    async def lose_one_label(*args):
        output = await original(*args)
        if args[:2] == ("pdftotext", "-bbox"):
            tree = ET.parse(args[-1])
            label = next(
                word
                for word in tree.iter("{http://www.w3.org/1999/xhtml}word")
                if word.text == "名称"
            )
            label.text = ""
            tree.write(args[-1], encoding="utf-8", xml_declaration=True)
        return output

    monkeypatch.setattr(module, "_process", lose_one_label)
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(), WeekDisplay("小一班", "第一学期", 3)
    )
    assert result.reason == "layout_overflow"
    assert not result.fits and result.data is None


@pytest.mark.parametrize("axis", ["x", "y"])
@pytest.mark.real_render
@pytest.mark.real_render
async def test_renderer_text_crossing_cell_edge_is_not_delivered(monkeypatch, axis):
    import xml.etree.ElementTree as ET

    from app.integration.word_export import shared_weekly_word as module

    original = module._process

    async def cross_cell(*args):
        output = await original(*args)
        if args[:2] == ("pdftotext", "-bbox"):
            tree = ET.parse(args[-1])
            word = next(
                word
                for word in tree.iter("{http://www.w3.org/1999/xhtml}word")
                if word.text == "项"
            )
            # Extend past the first date column's right edge, while its center
            # stays inside that original column and the bbox stays on the page.
            word.set("xMax", "240") if axis == "x" else word.set("yMax", "165")
            tree.write(args[-1], encoding="utf-8", xml_declaration=True)
        return output

    monkeypatch.setattr(module, "_process", cross_cell)
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(6), WeekDisplay("小一班", "第一学期", 3)
    )
    assert not result.fits, "text crosses a cell border but is delivered"
    assert result.reason == "layout_overflow" and result.data is None


@pytest.mark.parametrize("count", [5, 6])
def test_reference_layout_preserves_title_and_usable_content_width(count):
    doc = Document(
        BytesIO(
            fill_document(
                SEED_PATH.read_bytes(),
                complete_body(count),
                WeekDisplay("小一班", "第一学期", 3),
            )
        )
    )
    table = doc.tables[0]
    assert doc.paragraphs[0].runs[0].font.size.pt == 16
    assert doc.paragraphs[0].alignment == 1
    assert doc.paragraphs[0].runs[0].bold is True
    assert doc.paragraphs[0].runs[0].font.cs_bold is True
    assert table.columns[0].width < table.columns[2].width
    printable = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    assert abs(sum(c.width for c in table.columns) - printable) < 635 * (count + 2)
    for i in range(2, count + 2):
        assert table.cell(0, i).width == table.columns[i].width
    for row in range(5, 9):
        assert table.cell(row, 0)._tc is not table.cell(row, 1)._tc
        assert table.cell(row, 1)._tc is table.cell(row, count + 1)._tc
    assert "集体游戏：1." in table.cell(3, 2).text
    assert "自主游戏：" in table.cell(3, 2).text
    assert "本周重点指导区域：" in table.cell(4, 2).text
    assert "\n2." in table.cell(5, 1).text


@pytest.mark.parametrize("count", [5, 6])
@pytest.mark.parametrize("start", [date(2026, 9, 28), date(2026, 12, 28)])
def test_week_dates_follow_class_and_table_header_is_weekdays_only(count, start):
    body = complete_body(count)
    days = tuple(
        replace(d, day=start + timedelta(days=i)) for i, d in enumerate(body.days)
    )
    body = replace(
        body,
        base=replace(body.base, days=days),
        calendar=replace(
            body.calendar,
            columns=tuple(
                (d.day, teaching, label)
                for d, (_, teaching, label) in zip(days, body.calendar.columns)
            ),
        ),
    )
    doc = Document(
        BytesIO(
            fill_document(
                SEED_PATH.read_bytes(), body, WeekDisplay("中四班", "第一学期", 3)
            )
        )
    )
    first, last = days[0].day, days[-1].day
    end_year = f"{last.year}年" if last.year != first.year else ""
    expected = f"班级：中四班 第3周（{first.year}年{first.month}月{first.day}日—{end_year}{last.month}月{last.day}日）"
    assert expected in doc.paragraphs[1].text
    table = doc.tables[0]
    assert table.cell(0, 0).text == table.cell(0, 1).text == ""
    assert [table.cell(0, i + 2).text for i in range(count)] == [
        "周" + "一二三四五六日"[d.day.weekday()] for d in days
    ]


TITLE_STROKE_FIXTURE = (
    Path(__file__).resolve().parents[1] / "tests/fixtures/wpe-title-stroke-20260913.svg"
)


def _svg_namespace(tag: str) -> str:
    return "{http://www.w3.org/2000/svg}" + tag


def _load_frozen_title_stroke_svg():
    return ET.parse(TITLE_STROKE_FIXTURE)


def _stroked_paths(svg):
    return [
        path
        for path in svg.iter(_svg_namespace("path"))
        if path.get("fill") == "none" and path.get("stroke") is not None
    ]


def _frozen_svg_with_diagnostic(d: str, transform: str = "matrix(1, 0, 0, 1, 0, 0)"):
    """Copy the known-good frozen SVG plus ONE diagnostic stroked path."""
    svg = _load_frozen_title_stroke_svg()
    path = ET.Element(_svg_namespace("path"))
    path.set("fill", "none")
    path.set("stroke", "rgb(0%, 0%, 0%)")
    path.set("d", d)
    path.set("transform", transform)
    svg.getroot().append(path)
    return svg


def test_frozen_native_title_stroke_svg_parses_real_table_axes():
    svg = _load_frozen_title_stroke_svg()
    assert len(_stroked_paths(svg)) == 19
    vertical, horizontal = _rendered_grid(svg)
    assert len(vertical) == 8
    assert len(horizontal) == 10
    assert min(vertical) == pytest.approx(35.957, abs=0.01)
    assert max(vertical) == pytest.approx(558.648, abs=0.01)
    assert min(horizontal) == pytest.approx(74.801, abs=0.01)
    assert max(horizontal) == pytest.approx(602.582, abs=0.01)


def _native_title_contour() -> str:
    return next(
        path.get("d")
        for path in _stroked_paths(_load_frozen_title_stroke_svg())
        if "C" in path.get("d", "")
    )


def test_native_trailing_complete_move_only_subpath_is_accepted():
    """The actual contour ends 'Z M x y' (complete trailing move); it parses.

    With the identity transform the native coordinates lie below the real
    table bounds, so the frozen grid still resolves 8/10 clustered axes —
    success proves the parser accepts the native trailing move-only subpath.
    """
    svg = _frozen_svg_with_diagnostic(_native_title_contour())
    vertical, horizontal = _rendered_grid(svg)
    assert len(vertical) == 8 and len(horizontal) == 10


def test_supported_closed_contour_outside_table_ending_before_z_is_accepted():
    """A supported C/Z contour outside the table may end at a point different
    from its start: Z implicitly closes the drawn subpath, and the trailing
    move-only subpath completes the contour."""
    outside = "M 20 760 C 30 750 40 750 50 760 C 60 770 30 770 20 765 Z M 50 780"
    svg = _frozen_svg_with_diagnostic(outside)
    vertical, horizontal = _rendered_grid(svg)
    assert len(vertical) == 8 and len(horizontal) == 10


def test_complete_move_only_subpath_outside_table_is_accepted():
    svg = _frozen_svg_with_diagnostic("M 10 10 Z M 20 20")
    vertical, horizontal = _rendered_grid(svg)
    assert len(vertical) == 8 and len(horizontal) == 10


def test_intersecting_supported_closed_contour_is_rejected():
    """A supported C/Z contour fully inside the real table bounds is rejected."""
    inside = "M 100 300 C 150 250 250 250 300 300 C 350 350 150 350 100 300 Z M 300 300"
    svg = _frozen_svg_with_diagnostic(inside)
    with pytest.raises(LayoutRejected, match="layout_geometry_invalid"):
        _rendered_grid(svg)


def test_contour_touching_table_edge_is_rejected():
    """A supported contour that touches (not only crosses) the table bounds
    must fail closed: the bounds already carry the strict-disjointness
    epsilon, so even exact contact is intersection."""
    # Real frozen table bounds: x 35.957..558.648, y 74.801..602.582.
    touching_left_edge = (
        "M 35.957 300 C 30 250 20 250 10 300 C 0 350 30 350 35.957 300 Z"
    )
    svg = _frozen_svg_with_diagnostic(touching_left_edge)
    with pytest.raises(LayoutRejected, match="layout_geometry_invalid"):
        _rendered_grid(svg)


def test_unsupported_contour_outside_table_is_rejected_geometry_invalid():
    svg = _frozen_svg_with_diagnostic("M 10 10 Q 20 5 30 10 Z")
    with pytest.raises(LayoutRejected, match="layout_geometry_invalid"):
        _rendered_grid(svg)


@pytest.mark.parametrize(
    ("d", "transform"),
    [
        # command without needed coords, then transition discards nothing
        ("M 10 10 C 12 12 Z", None),
        ("M 10 10 L 20", None),
        # incomplete parameters discarded by Z transition
        ("M 10 10 C 12 12 14 14 Z", None),
        # zero-parameter commands must be rejected
        ("M 10 10 C Z", None),
        ("M 10 10 M", None),
        ("M", None),
        # implicit extra parameters restart the completed command
        ("M 10 10 20 20", None),
        ("M 10 10 L 20 20 30 30 Z", None),
        # commands after Z except M
        ("M 10 10 Z L 20 20", None),
        # unclosed drawn subpath
        ("M 10 10 L 20 20", None),
        ("M 10 10 C 20 20 30 30 40 40", None),
        # relative command letter
        ("m 10 10 l 20 10 z", None),
        # unsupported quadratic command
        ("M 10 10 Q 22 8 24 10 Z", None),
        # Python-only float spelling (underscore) is not a valid SVG lexeme
        ("M 10 10 L 1_0 20 Z", None),
        # non-finite and junk tokens
        ("M nan nan L 20 10 Z", None),
        ("M 10 10 L 20 10 junk", None),
        # multiple unclosed subpaths
        ("M 10 10 L 20 20 M 30 30 L 40 40", None),
        # empty path
        ("", None),
    ],
)
def test_malformed_or_nonfinite_paths_are_rejected(d, transform):
    matrix = transform if transform is not None else "matrix(1, 0, 0, 1, 0, 0)"
    svg = _frozen_svg_with_diagnostic(d, matrix)
    with pytest.raises(LayoutRejected, match="layout_geometry_invalid"):
        _rendered_grid(svg)


@pytest.mark.parametrize("transform", ["matrix(1, 0, 0)", "matrix(inf, 0, 0, 1, 0, 0)"])
def test_malformed_transform_matrix_is_rejected(transform):
    svg = _frozen_svg_with_diagnostic("M 10 10 L 20 10 Z", transform)
    with pytest.raises(LayoutRejected, match="layout_geometry_invalid"):
        _rendered_grid(svg)


@pytest.mark.parametrize("count", [5, 6])
@pytest.mark.real_render
@pytest.mark.real_render
async def test_removing_a_required_border_keeps_actual_render_undelivered(
    count, monkeypatch
):
    from app.integration.word_export import shared_weekly_word as module

    original = module._process
    straight = r"M\s+(-?[0-9.]+)\s+(-?[0-9.]+)\s+L\s+(-?[0-9.]+)\s+(-?[0-9.]+)\s*"

    async def drop_one_border(*args):
        output = await original(*args)
        if args[0] == "pdftocairo":
            tree = ET.parse(args[-1])
            borders = [
                path
                for path in tree.iter(_svg_namespace("path"))
                if path.get("fill") == "none"
                and path.get("stroke") is not None
                and "C" not in path.get("d", "")
            ]
            assert len(borders) == count + 13, "real straight border path count"
            # The unique outer right edge is a required vertical axis source;
            # find it by parsed straight coordinates, not a literal guess.
            straight_borders = []
            for path in borders:
                match = re.fullmatch(straight, path.get("d", ""))
                assert match is not None
                x0, y0, x1, y1 = map(float, match.groups())
                straight_borders.append((path, x0, y0, x1, y1))
            verticals = [
                (path, x0, y0, y1)
                for path, x0, y0, x1, y1 in straight_borders
                if abs(x0 - x1) < 0.1 and abs(y1 - y0) > 5
            ]
            assert verticals
            target = max(verticals, key=lambda item: item[1])[0]
            parent = next(element for element in tree.iter() if target in list(element))
            parent.remove(target)
            remaining = [
                path
                for path in tree.iter(_svg_namespace("path"))
                if path.get("fill") == "none"
                and path.get("stroke") is not None
                and "C" not in path.get("d", "")
            ]
            assert len(remaining) == count + 12, "required border path must be removed"
            tree.write(args[-1], encoding="utf-8", xml_declaration=True)
        return output

    monkeypatch.setattr(module, "_process", drop_one_border)
    result = await SharedWeeklyWordPort().inspect_local(
        complete_body(count), WeekDisplay("小一班", "第一学期", 3)
    )
    assert not result.fits
    assert result.reason == "layout_overflow" and result.data is None
